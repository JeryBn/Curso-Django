from pathlib import Path
from copy import deepcopy
import hashlib
import json
import zipfile
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = Path(r'C:\Users\Chory\Downloads\GLAB-S06-MMONTGOMERY-2026-02.docx')
output = root/'GLAB-S06-MMONTGOMERY-2026-02-Jery-Becerra-RESUELTO.docx'
doc = Document(source)
original_parts = {name: hashlib.sha256(data).hexdigest() for name,data in [(n,zipfile.ZipFile(source).read(n)) for n in zipfile.ZipFile(source).namelist()]}
conclusion_index = next(i for i,p in enumerate(doc.paragraphs) if 'Conclusiones:' in p.text)
for para in list(doc.paragraphs)[conclusion_index:]:
    para._element.getparent().remove(para._element)
for name in ['Heading 1','Heading 2']:
    style=doc.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0); style.font.size=Pt(15 if name=='Heading 1' else 12)
    style.paragraph_format.space_before=Pt(8); style.paragraph_format.space_after=Pt(8)
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)

def p(text):
    para=doc.add_paragraph(text)
    para.paragraph_format.space_after=Pt(7)
    para.paragraph_format.line_spacing=1.05
    for run in para.runs: run.font.name='Arial';run.font.size=Pt(11);run.font.color.rgb=RGBColor(0,0,0)
    return para
def page(title): doc.add_page_break();doc.add_heading(title,1)
def code(text):
    for line in text.strip().splitlines():
        para=doc.add_paragraph();para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1
        run=para.add_run(line);run.font.name='Consolas';run.font.size=Pt(9)
    p('')
def picture(name,caption,max_height=3.7):
    file=root/'evidencias'/name
    w,h=Image.open(file).size
    width=min(6.25,max_height*w/h)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(file),width=Inches(width))
    para.paragraph_format.space_after=Pt(4)
    p(caption)
def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=text
    for row_index,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            if row_index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF2');pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
                for run in para.runs:run.font.name='Arial';run.font.size=Pt(10);run.bold=row_index==0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

page('Desarrollo del laboratorio 6')
p('Alumno responsable: Jery Becerra Ninaquispe. Fecha de elaboración y verificación: 8 de octubre de 2026.')
p('Se implementó Actualidad Digital, un portal con seis noticias en tres categorías. La aplicación news continúa el proyecto de la semana 5, conserva movies y reutiliza su modelo abstracto de auditoría. La portada, el detalle y el listado por categoría utilizan una base común y fragmentos reutilizables; el contenido se administra desde Django.')
p('Responsabilidad individual: Jery Becerra Ninaquispe presenta la entrega. Se utilizó asistencia de Codex para implementación, validación y documentación, registrada en PROMPTS.md. La guía original se conserva como referencia y esta copia incorpora la resolución y las evidencias.')
table(['Pasos de la guía','Implementación y evidencia'],[
('1 a 3','Aplicación news, Pillow, configuración, modelos, migraciones y superusuario local.'),
('4 y 5','base.html con tres bloques y _article_card.html reutilizado.'),
('6','Portada con for, empty, date y truncatewords.'),
('7 y 8','Detalle con autor e imagen y listado filtrado por categoría.'),
('9 y 10','Rutas con nombres, etiqueta url, static y archivos de medios.'),
('11','Tres ModelAdmin, seis noticias, tres categorías y cambio visible desde el panel.'),
('12','HTML guardado en la noticia y mostrado como texto literal.'),
('13','Código y documentación publicados en Curso-Django, carpeta Semana 6.')],[1.2,5.05])
p('Ubicación local: '+str(root))
p('Portal: http://127.0.0.1:8096/ · Administrador: http://127.0.0.1:8096/admin/')

page('Configuración y relaciones entre modelos')
p('Pasos 1 a 3. Se instaló Django 5.2.17 y Pillow 12.3.0 en un entorno virtual propio. news se agregó a INSTALLED_APPS; templates es el directorio compartido y los recursos de noticias se organizan en static/news y media/articles.')
code('''INSTALLED_APPS = [
    # Django applications and the previous movies application
    "news.apps.NewsConfig",
]
TEMPLATES[0]["DIRS"] = [BASE_DIR / "templates"]
STATICFILES_DIRS = [BASE_DIR / "static"]
MEDIA_ROOT = BASE_DIR / "media"
MEDIA_URL = "media/"
if settings.DEBUG:
    urlpatterns += static(settings.MEDIA_URL,
                          document_root=settings.MEDIA_ROOT)''')
p('Article guarda título, slug, resumen, cuerpo, imagen destacada, fecha y estado de publicación. Su autor es una relación ForeignKey con protección frente a eliminación; las categorías utilizan ManyToManyField. Author tiene nombre, biografía y fotografía; Category tiene nombre, slug y descripción. Las tres entidades heredan created_at y updated_at.')
code('''author = models.ForeignKey(
    Author, on_delete=models.PROTECT, related_name="articles")
categories = models.ManyToManyField(
    Category, related_name="articles")
featured_image = models.ImageField(
    upload_to="articles/", blank=True)
published_at = models.DateTimeField(default=timezone.now)''')
p('Se ejecutaron makemigrations news y migrate. La migración inicial de news quedó versionada. Se creó admin_lab como superusuario local, sin publicar su contraseña. seed_news genera los seis artículos, las categorías y las imágenes para reproducir la entrega sin copiar la base de datos.')
p('Caso de prueba: manage.py check no detectó problemas y makemigrations --check --dry-run informó No changes detected. Las relaciones y la selección de artículos se verificaron en las pruebas automatizadas.')

page('Herencia y estructura de las plantillas')
p('Pasos 4 y 5. base.html define el encabezado, la navegación, el contenido, la barra lateral y el pie. Las tres páginas heredan esa estructura. El fragmento de tarjeta se incluye desde la portada y la categoría, evitando repetir el marcado de una noticia.')
picture('10_editor_estructura.png','Figura 1. Estructura del proyecto y home.html en Visual Studio Code. Se observan extends, for, include y empty.',3.65)
code('''{% extends 'news/base.html' %}
{% block title %}Portada de Actualidad Digital{% endblock %}
{% block content %}
  {% for article in articles %}
    {% include 'news/_article_card.html' with article=article only %}
  {% empty %}
    <h2>Aún no hay noticias publicadas</h2>
  {% endfor %}
{% endblock %}''')
p('El bloque sidebar incluye _sidebar.html desde la base. Las pruebas comprueban el uso efectivo de base.html, _article_card.html y _sidebar.html al renderizar la portada; la categoría utiliza el mismo fragmento de tarjeta.')

page('Portada con datos del modelo')
p('Paso 6. La vista consulta las noticias publicadas y ordenadas por fecha. La plantilla recorre articles y cada tarjeta presenta el título, autor, fecha, imagen, resumen y categorías. El resumen se recorta con truncatewords y la fecha se formatea con date.')
picture('11_portada_escritorio.jpg','Figura 2. Portada alimentada por la base de datos. El título de la primera noticia ya refleja el cambio realizado desde el administrador.',3.9)
code('''{{ article.title }}
{{ article.author.name }}
{{ article.published_at|date:"d M Y" }}
{{ article.summary|truncatewords:22 }}
{% if article.featured_image %}
  <img src="{{ article.featured_image.url }}" alt="...">
{% endif %}''')
p('La lógica para decidir qué noticias pueden publicarse está en ArticleQuerySet.published y en las vistas. Las plantillas presentan los datos mediante variables, filtros y etiquetas de control; no calculan reglas de negocio.')

page('Listado de noticias por categoría')
p('Paso 8. La ruta recibe el slug de una categoría. La vista obtiene la categoría y filtra sus noticias publicadas. category.html hereda de base.html e incluye _article_card.html dentro de su propio recorrido.')
picture('12_categoria_escritorio.jpg','Figura 3. Tecnología muestra las dos noticias asociadas a esa categoría y conserva la misma tarjeta y navegación de la portada.',3.9)
code('''selected_category = get_object_or_404(Category, slug=slug)
articles = Article.objects.published().filter(
    categories=selected_category)
{% for article in articles %}
  {% include 'news/_article_card.html' with article=article only %}
{% empty %}
  <h2>Sin noticias en esta categoría</h2>
{% endfor %}''')
p('Caso de prueba: el listado muestra el artículo de su categoría y excluye uno que no pertenece a ella. También se comprueba el mensaje sin resultados y la respuesta 404 para una categoría inexistente.')

page('Detalle y navegación entre páginas')
p('Pasos 7 y 9. detail.html hereda de la base y muestra el artículo seleccionado con imagen, autor, fecha, categorías, resumen y cuerpo. Las rutas tienen los nombres news:home, news:detail y news:category; los enlaces utilizan la etiqueta url.')
picture('13_detalle_escritorio.jpg','Figura 4. Detalle de la noticia después de modificar su título en el administrador.',3.8)
code('''path("", views.home, name="home")
path("noticia/<slug:slug>/", views.detail, name="detail")
path("categoria/<slug:slug>/", views.category, name="category")
<a href="{% url 'news:detail' article.slug %}">Leer noticia</a>
<a href="{% url 'news:category' category.slug %}">
  {{ category.name }}
</a>''')
p('La navegación se comprobó desde portada a categoría y desde la tarjeta al detalle. Las pruebas verifican autor, categoría y fecha en el detalle y devuelven 404 para noticias inexistentes, borradores o publicaciones futuras.')

page('Archivos estáticos e imágenes')
p('Paso 10. base.html carga static y obtiene la hoja portal.css mediante la etiqueta static. Las imágenes son archivos PNG generados con Pillow y guardados en ImageField; config/urls.py sirve MEDIA_URL durante desarrollo.')
code('''{% load static %}
<link rel="stylesheet" href="{% static 'news/portal.css' %}">
{% if article.featured_image %}
  <img src="{{ article.featured_image.url }}"
       alt="Imagen de {{ article.title }}">
{% else %}
  <div class="image-placeholder">Actualidad Digital</div>
{% endif %}''')
p('La hoja de estilos se localizó mediante el sistema staticfiles; la imagen de la primera noticia existe en media/articles. Las capturas del portal muestran el CSS aplicado y las imágenes cargadas. La imagen es opcional: si no existe, la tarjeta presenta una alternativa sin intentar acceder a una URL vacía.')
p('Observación: el servidor de desarrollo y DEBUG=True se usan únicamente para este laboratorio local. La publicación en GitHub conserva el código; no constituye un despliegue del servidor Django en Internet.')
picture('01_portada.jpg','Figura 5. Vista estrecha del mismo portal. El diseño adapta la navegación y las tarjetas al ancho disponible.',3.5)

page('Contenido gestionado desde el administrador')
p('Paso 11. Se configuraron ArticleAdmin, CategoryAdmin y AuthorAdmin con list_display, list_filter y search_fields. ArticleAdmin añade selección horizontal de categorías y navegación por fecha. El administrador permite gestionar las tres entidades sin modificar las plantillas.')
picture('04_admin_listado.jpg','Figura 6. El listado del administrador contiene seis noticias, búsqueda, columnas y filtros.',3.9)
code('''list_display = ("title", "author", "published_at", "is_published")
list_filter = ("is_published", "categories", "author", "published_at")
search_fields = ("title", "summary", "body", "author__name")
filter_horizontal = ("categories",)
date_hierarchy = "published_at"''')
p('Los datos iniciales se cargaron con seed_news y se comprobaron en el panel: seis noticias y tres categorías. Después se modificaron realmente el título y el cuerpo de la primera noticia desde el formulario del administrador. El cambio se conservó en la base de datos y se vio en el sitio público.')

page('Edición del panel publicada en el sitio')
p('La noticia tenía el título Django conecta el contenido con sus plantillas. Desde el administrador se cambió a Django publica cambios desde el administrador y se agregó un nuevo cuerpo con la prueba HTML. Se guardó el formulario y se recargó el detalle público.')
picture('15_admin_formulario.jpg','Figura 7. Formulario de la noticia con el título y el cuerpo modificados. La figura 4 muestra el mismo título en el portal.',3.9)
p('Resultado: el sitio muestra los valores guardados sin editar home.html, detail.html ni el fragmento de tarjeta. Esta comprobación satisface el criterio de gestionar contenido desde el panel y publicarlo en las plantillas.')
p('Caso automatizado: test_admin_edit_is_visible_without_code_changes envía un cambio al formulario real del administrador, comprueba la redirección de guardado y confirma que el título actualizado aparece en la portada.')
code('''response = self.client.post(admin_change_url, article_data)
self.assertEqual(response.status_code, 302)
self.assertContains(self.client.get(reverse("news:home")),
                    "Updated from admin")''')

page('Prueba de escapado automático de HTML')
p('Paso 12. Se guardó en el cuerpo la cadena <strong>Texto importante</strong> y <script>alert("prueba")</script>. La página muestra los caracteres de las etiquetas como texto: no convierte strong en negrita ni ejecuta el script.')
picture('14_escapado_escritorio.jpg','Figura 8. El cuerpo público presenta literalmente las etiquetas HTML guardadas desde el administrador.',3.8)
code('''<div class="body-text">{{ article.body|linebreaks }}</div>
# Expected content in the HTTP response
&lt;strong&gt;Texto importante&lt;/strong&gt;
&lt;script&gt;alert(&quot;prueba&quot;)&lt;/script&gt;''')
p('Django escapa automáticamente el contenido de la variable. linebreaks conserva el escapado y agrega párrafos para los saltos de línea. No se utiliza safe ni autoescape off. El texto del artículo sigue siendo contenido y no se interpreta como código del sitio.')
p('Caso de prueba: test_html_is_escaped verifica las entidades &lt;strong&gt; y &lt;script&gt; en la respuesta y comprueba que no aparece un bloque script sin escapar.')

page('Caso sin noticias y pruebas realizadas')
p('Para verificar empty se ocultaron temporalmente las seis noticias, se recargó la portada y se capturó el mensaje de ausencia de contenido. Luego se restauró su estado publicado: la entrega final conserva las seis noticias visibles.')
picture('17_empty_escritorio.jpg','Figura 9. Mensaje del bloque empty cuando no existen noticias publicadas. La operación fue temporal y los datos se restauraron.',3.55)
table(['Caso','Resultado comprobado'],[
('Herencia y fragmentos','La portada usa base, tarjeta y barra lateral.'),
('Recorrido y filtros','Fecha formateada y resumen recortado.'),
('Categoría y empty','Filtra correctamente y muestra el caso vacío.'),
('Detalle y 404','Autor y categorías visibles; rutas inválidas dan 404.'),
('Publicación y escapado','Oculta borradores y fechas futuras; HTML escapado.'),
('Cambio desde admin','Título guardado visible en el sitio sin cambiar código.')],[1.5,4.75])
p('Resultado final: 25 pruebas aprobadas, 10 de news y 15 de movies. El registro completo está en evidencias/pruebas.txt. manage.py check no reportó problemas y no quedaron migraciones pendientes.')

page('Repositorio y observaciones de entrega')
p('Paso 13. Se publicó el proyecto en el repositorio público JeryBn/Curso-Django, rama main, carpeta existente Semana 6. Se conservó ese nombre para mantener la estructura del curso. La subida de código fue verificada en GitHub con el commit c3fd5bb82f4725ed8fb023d065902eadc6662acd.')
picture('16_github_publicacion.jpg','Figura 10. Semana 6 publicada en GitHub con news, plantillas, archivos estáticos, evidencias y documentación.',3.8)
p('Enlace de revisión: https://github.com/JeryBn/Curso-Django/tree/main/Semana%206')
p('README.md explica la instalación, la estructura de plantillas, sus responsabilidades y las observaciones del laboratorio. EVIDENCIAS.md identifica las capturas y PROMPTS.md registra la asistencia. La guía resuelta se incorpora a la misma carpeta como parte de la entrega final.')
p('Los archivos db.sqlite3, .venv y media no se versionan. seed_news reproduce los datos y genera las imágenes. Las contraseñas y los tokens de GitHub no forman parte del código ni del Word. La presentación al campus virtual debe realizarse desde la cuenta del alumno.')

page('Conclusiones del laboratorio')
for text in [
'La herencia centraliza el encabezado, la navegación y la distribución del portal en base.html. Los bloques permiten que cada página aporte su contenido sin copiar la estructura común.',
'El fragmento _article_card.html se reutiliza desde la portada y la categoría. Una modificación de la tarjeta se aplica a ambos listados y reduce el riesgo de diferencias entre páginas.',
'Las variables muestran los datos del modelo y las etiquetas for, empty e if controlan su presentación. date, truncatewords y default formatean los valores sin trasladar reglas de publicación a las plantillas.',
'El administrador gestiona Article, Author y Category. La edición real de una noticia demostró que el contenido actualizado se publica sin modificar el código de las plantillas.',
'El escapado automático mantiene el HTML almacenado como texto. La prueba con strong y script confirmó que la página no interpreta etiquetas aportadas por el contenido de la noticia.',
'La entrega conserva la base de la semana 5 y agrega news por responsabilidad. Las migraciones, el comando de datos, las pruebas, las observaciones y las evidencias permiten reproducir y revisar el resultado publicado.'
]:p(text)
table(['Criterio de rúbrica','Evidencia incluida'],[
('Herencia y fragmentos','Base, tres páginas heredadas, tarjeta y barra lateral reutilizadas.'),
('Datos y filtros','Portada, categorías, detalle y empty con variables y filtros.'),
('Gestión desde admin','Tres ModelAdmin y cambio real publicado.'),
('Repositorio y observaciones','Publicación remota y README con estructura explicada.')],[2.0,4.25])
p('Las cuatro áreas de la rúbrica tienen implementación y evidencia. La calificación final corresponde al docente.')
doc.save(output)
with zipfile.ZipFile(output) as final:
    preserved=[name for name,h in original_parts.items() if name in final.namelist() and hashlib.sha256(final.read(name)).hexdigest()==h]
contract={'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'reference_pages':4,'sections':len(doc.sections),'page_system':[[s.page_width/914400,s.page_height/914400,s.top_margin/914400,s.bottom_margin/914400,s.left_margin/914400,s.right_margin/914400] for s in doc.sections],'preserved_parts':preserved,'changes':'Original cover, requirements, procedure and rubric retained. Blank conclusions replaced by appended development and conclusions. Added paragraphs use Arial and retain original page geometry and headers.'}
(root/'_qa/artifact.md').write_text(json.dumps(contract,indent=2,ensure_ascii=False),encoding='utf-8')
print(output)
print('Preserved package parts:',len(preserved),'of',len(original_parts))

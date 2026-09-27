from __future__ import annotations

import argparse
import textwrap
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BLUE = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF4FB"
GREEN = "1F6D48"
LIGHT_GREEN = "E2F0D9"
GRAY = "F2F2F2"
MID_GRAY = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_field(paragraph) -> None:
    paragraph.add_run("Página ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def ensure_styles(document: Document) -> None:
    styles = document.styles

    def create_style(name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
        return styles[name] if name in styles else styles.add_style(name, style_type)

    body = create_style("GLAB Body")
    body.font.name = "Calibri"
    body.font.size = Pt(10.5)
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.08

    title = create_style("GLAB Annex Title")
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLUE)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    heading1 = create_style("GLAB Heading 1")
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor.from_string(BLUE)
    heading1.paragraph_format.space_before = Pt(10)
    heading1.paragraph_format.space_after = Pt(6)
    heading1.paragraph_format.keep_with_next = True

    heading2 = create_style("GLAB Heading 2")
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(12.5)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor.from_string(GREEN)
    heading2.paragraph_format.space_before = Pt(8)
    heading2.paragraph_format.space_after = Pt(5)
    heading2.paragraph_format.keep_with_next = True

    caption = create_style("GLAB Caption")
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    note = create_style("GLAB Note")
    note.font.name = "Calibri"
    note.font.size = Pt(9.5)
    note.font.color.rgb = RGBColor.from_string(BLUE)
    note.paragraph_format.left_indent = Cm(0.4)
    note.paragraph_format.right_indent = Cm(0.4)
    note.paragraph_format.space_after = Pt(8)


def add_body(document: Document, text: str, bold_lead: str | None = None):
    paragraph = document.add_paragraph(style="GLAB Body")
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="GLAB Body")
    paragraph.style = "GLAB Body"
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    paragraph.add_run("• ").bold = True
    paragraph.add_run(text)


def add_status_box(document: Document, title: str, text: str, fill=LIGHT_GREEN) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.7)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    lead = paragraph.add_run(f"{title}: ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(GREEN)
    paragraph.add_run(text)
    document.add_paragraph(style="GLAB Body").paragraph_format.space_after = Pt(1)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cell = cells[index]
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Calibri"
            run.font.size = Pt(8.7)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph(style="GLAB Body").paragraph_format.space_after = Pt(1)


def add_code_block(document: Document, title: str, code: str) -> None:
    document.add_paragraph(title, style="GLAB Heading 2")
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, GRAY)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(textwrap.dedent(code).strip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(7.8)
    document.add_paragraph(style="GLAB Body").paragraph_format.space_after = Pt(1)


def add_evidence_image(
    document: Document,
    number: int,
    title: str,
    explanation: str,
    image_path: Path,
) -> None:
    heading = document.add_paragraph(f"Evidencia {number}. {title}", style="GLAB Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.left_indent = Cm(0)
    heading.paragraph_format.right_indent = Cm(0)
    heading.paragraph_format.first_line_indent = Cm(0)
    add_body(document, explanation)
    max_width = 6.35
    max_height = 6.65
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width_in = min(max_width, max_height * ratio)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_in))
    caption = document.add_paragraph(
        f"Figura {number}. Captura obtenida durante la verificación local del laboratorio.",
        style="GLAB Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_new_section(document: Document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)

    # La plantilla usa encabezados distintos en páginas pares. Se limpian todas
    # las variantes para evitar que formas heredadas se superpongan al anexo.
    for header_part in (
        section.header,
        section.even_page_header,
        section.first_page_header,
    ):
        header_part.is_linked_to_previous = False
        for child in list(header_part._element):
            header_part._element.remove(child)
        header = header_part.add_paragraph()
        header.text = "GLAB-S05  |  Administración avanzada con Django"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    for footer_part in (
        section.footer,
        section.even_page_footer,
        section.first_page_footer,
    ):
        footer_part.is_linked_to_previous = False
        for child in list(footer_part._element):
            footer_part._element.remove(child)
        footer = footer_part.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("Michael Montgomery Rosell  •  ")
        footer_run.font.name = "Calibri"
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor.from_string(MID_GRAY)
        add_page_field(footer)
        for run in footer.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    return section


def build_document(source: Path, destination: Path, project_root: Path) -> None:
    document = Document(str(source))
    ensure_styles(document)

    # Completa la sección de conclusiones existente sin modificar la guía original.
    conclusion_heading = None
    original_paragraphs = document.paragraphs
    heading_index = None
    for index in range(len(original_paragraphs) - 1, -1, -1):
        paragraph = original_paragraphs[index]
        if paragraph.text.strip().lower().startswith("conclusiones"):
            conclusion_heading = paragraph
            heading_index = index
            break
    if conclusion_heading is not None and heading_index is not None:
        for candidate in original_paragraphs[max(0, heading_index - 3) : heading_index]:
            if candidate.text.strip() == "1.":
                candidate._element.getparent().remove(candidate._element)
        for candidate in original_paragraphs[heading_index + 1 : heading_index + 4]:
            if not candidate.text.strip():
                candidate._element.getparent().remove(candidate._element)
        conclusion_heading.text = "Conclusiones:"

    add_body(
        document,
        "La práctica permitió comprobar que el administrador de Django ofrece una base CRUD funcional y que su valor aumenta cuando cada modelo dispone de columnas, filtros, búsqueda y relaciones adaptadas al dominio. Los cuatro modelos quedaron registrados y pueden administrarse desde una interfaz coherente en español.",
    )
    add_body(
        document,
        "Las relaciones entre películas, géneros, personas y valoraciones se implementaron con ForeignKey y ManyToManyField. El bloque de valoraciones en línea reduce pasos durante la edición de una película, mientras que los campos de creación y modificación aportan trazabilidad sin permitir cambios manuales.",
    )
    add_body(
        document,
        "La comparación entre el superusuario y el usuario editor demostró que Django aplica los permisos tanto en la navegación como en las operaciones. El grupo editores puede ver, agregar y modificar películas, pero no eliminarlas ni administrar usuarios o grupos.",
    )
    add_body(
        document,
        "La vista pública completa el caso de uso al recomendar películas que comparten género y ordenar los resultados por el promedio de las valoraciones. Las quince pruebas automatizadas confirmaron las reglas de integridad, la configuración de ModelAdmin, las restricciones de acceso y el algoritmo de recomendación.",
    )

    configure_new_section(document)

    title = document.add_paragraph("ANEXO DE DESARROLLO Y EVIDENCIAS", style="GLAB Annex Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(style="GLAB Body")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Implementación, verificación y evidencias del funcionamiento").bold = True
    subtitle.add_run("\nFecha de verificación: 26 de septiembre de 2026")

    add_status_box(
        document,
        "Resultado",
        "laboratorio implementado y verificado localmente: comprobación de Django sin observaciones, sin migraciones pendientes y 15 de 15 pruebas aprobadas.",
    )

    document.add_paragraph("1. Correspondencia con los objetivos de la guía", style="GLAB Heading 1")
    add_table(
        document,
        ["Requisito", "Implementación realizada", "Comprobación"],
        [
            ["Modelos relacionados", "Genre, Person, Movie y Rating; relaciones M:N y 1:N.", "Migración movies.0001 aplicada."],
            ["ModelAdmin", "Columnas, filtros, búsqueda, orden, autocompletado y selección horizontal.", "Listados y búsqueda probados en el panel."],
            ["Inline", "RatingInline dentro del formulario de Movie.", "Formulario con valoraciones en la misma pantalla."],
            ["Auditoría", "created_at y updated_at heredados y de solo lectura.", "Campos visibles sin controles editables."],
            ["Datos", "4 géneros, 8 personas, 10 películas y 15 valoraciones.", "Conteos revisados en la base local."],
            ["Acceso", "Superusuario, grupo editores y usuario editor restringido.", "El editor recibe 403 al intentar eliminar."],
            ["Recomendación", "Mismo género, con valoración, promedio descendente.", "Vista pública y pruebas automatizadas."],
        ],
        widths=[3.2, 8.1, 5.0],
    )

    document.add_paragraph("2. Entorno y estructura", style="GLAB Heading 1")
    add_table(
        document,
        ["Elemento", "Valor verificado"],
        [
            ["Python", "3.12.14"],
            ["Django", "5.2.17"],
            ["Pillow", "12.3.0"],
            ["Base de datos", "SQLite local"],
            ["Puerto", "8095"],
            ["Rutas", "/ para catálogo; /admin/ para administración"],
        ],
        widths=[5.0, 11.3],
    )
    add_code_block(
        document,
        "Estructura principal del proyecto",
        """
        config/                         configuración y rutas generales
        movies/
          admin.py                      ModelAdmin e inline
          models.py                     modelos y relaciones
          views.py                      catálogo y recomendaciones
          tests.py                      15 pruebas automatizadas
          management/commands/setup_lab.py
        templates/                      interfaz pública
        static/movies/site.css          estilos visuales
        evidencias/                     capturas numeradas
        manage.py                       punto de entrada
        """,
    )

    document.add_paragraph("3. Elaboración del código", style="GLAB Heading 1")
    add_body(
        document,
        "La implementación se organizó en capas. Los modelos expresan las relaciones y reglas de integridad; el administrador adapta cada operación al caso de uso; la vista pública calcula promedios mediante agregaciones; y el comando setup_lab reproduce usuarios, permisos y datos sin guardar contraseñas en el código.",
    )
    add_code_block(
        document,
        "3.1. Relaciones y auditoría en movies/models.py",
        """
        class TimeStampedModel(models.Model):
            created_at = models.DateTimeField("fecha de creación", auto_now_add=True)
            updated_at = models.DateTimeField("última modificación", auto_now=True)

            class Meta:
                abstract = True

        class Movie(TimeStampedModel):
            genres = models.ManyToManyField(Genre, related_name="movies")
            director = models.ForeignKey(
                Person, on_delete=models.PROTECT, related_name="directed_movies"
            )
            cast = models.ManyToManyField(
                Person, related_name="acting_movies", blank=True
            )

        class Rating(TimeStampedModel):
            movie = models.ForeignKey(
                Movie, on_delete=models.CASCADE, related_name="ratings"
            )
            score = models.PositiveSmallIntegerField(
                validators=[MinValueValidator(1), MaxValueValidator(5)]
            )
        """,
    )
    add_code_block(
        document,
        "3.2. Personalización e inline en movies/admin.py",
        """
        class RatingInline(admin.TabularInline):
            model = Rating
            extra = 1
            fields = ("reviewer_name", "score", "comment", "created_at", "updated_at")
            readonly_fields = ("created_at", "updated_at")

        @admin.register(Movie)
        class MovieAdmin(AuditFieldsAdminMixin, admin.ModelAdmin):
            list_display = (
                "title", "release_year", "genre_list", "director",
                "average_rating", "rating_count", "updated_at",
            )
            list_filter = ("genres", "release_year")
            search_fields = (
                "title", "synopsis", "director__first_name",
                "director__last_name", "genres__name",
            )
            filter_horizontal = ("genres", "cast")
            inlines = (RatingInline,)
        """,
    )
    add_code_block(
        document,
        "3.3. Algoritmo de recomendación en movies/views.py",
        """
        recommendations = (
            Movie.objects.annotate(
                average_rating=Avg("ratings__score"),
                rating_count=Count("ratings", distinct=True),
            )
            .filter(
                genres__in=selected_movie.genres.all(),
                rating_count__gt=0,
            )
            .exclude(pk=selected_movie.pk)
            .distinct()
            .order_by("-average_rating", "title")[:6]
        )
        """,
    )
    add_code_block(
        document,
        "3.4. Permisos reproducibles en setup_lab.py",
        """
        permission_names = {
            ("movie", "view_movie"),
            ("movie", "add_movie"),
            ("movie", "change_movie"),
            ("genre", "view_genre"),
            ("person", "view_person"),
            ("rating", "view_rating"),
        }
        group.permissions.set(selected)
        editor_user.is_staff = True
        editor_user.is_superuser = False
        editor_user.groups.set([group])
        """,
    )

    document.add_page_break()
    document.add_paragraph("4. Verificación automatizada", style="GLAB Heading 1")
    add_body(
        document,
        "Se ejecutaron las comprobaciones del framework, la detección de cambios de migraciones y el conjunto completo de pruebas. Los resultados corresponden a la ejecución final realizada en el proyecto entregado.",
    )
    add_code_block(
        document,
        "Resultado de la ejecución",
        """
        Python 3.12.14
        Django 5.2.17
        Pillow 12.3.0
        System check identified no issues (0 silenced).
        No changes detected
        Found 15 test(s).
        Ran 15 tests in 11.721s
        OK
        """,
    )
    add_table(
        document,
        ["Grupo de pruebas", "Validación principal", "Estado"],
        [
            ["ModelTests", "Representaciones y restricciones de unicidad.", "Aprobado"],
            ["AdminConfigurationTests", "Registro, filtros, búsqueda, inline y auditoría.", "Aprobado"],
            ["PermissionTests", "Permisos del editor y denegación de eliminación.", "Aprobado"],
            ["RecommendationViewTests", "Catálogo, género compartido, orden y exclusiones.", "Aprobado"],
        ],
        widths=[4.8, 8.9, 2.6],
    )
    add_status_box(
        document,
        "Prueba de seguridad funcional",
        "editor_lab puede ver, agregar y modificar películas, pero la ruta de eliminación responde con HTTP 403 y el panel no muestra Usuarios ni Grupos.",
        fill=LIGHT_BLUE,
    )

    evidence_dir = project_root / "evidencias"
    evidence_items = [
        (
            1,
            "Registro básico de los modelos",
            "Se registraron Genre, Movie, Person y Rating. Esta captura conserva el estado inicial del administrador y demuestra que Django generó el CRUD antes de aplicar la configuración avanzada.",
            "01_admin_registro_basico.jpg",
        ),
        (
            2,
            "Listado personalizado de películas",
            "El listado muestra título, año, géneros, director, promedio, número de valoraciones y modificación. A la derecha aparecen los filtros por género y año requeridos por la guía.",
            "02_admin_personalizado_listado.jpg",
        ),
        (
            3,
            "Búsqueda por título",
            "La búsqueda del término Marea devuelve únicamente la película coincidente. Esto comprueba el uso de search_fields en MovieAdmin.",
            "03_admin_busqueda_titulo.jpg",
        ),
        (
            4,
            "Valoraciones en línea y auditoría",
            "El formulario permite administrar valoraciones sin abandonar la película. También presenta las fechas de creación y modificación como información de solo lectura.",
            "04_admin_inline_auditoria_detalle.jpg",
        ),
        (
            5,
            "Grupo editores y permisos",
            "La selección visible incluye ver, agregar y cambiar películas, pero no eliminar. Los permisos auxiliares se limitan a consulta.",
            "05_grupo_editores_permisos.jpg",
        ),
        (
            6,
            "Panel restringido del usuario editor",
            "Al iniciar sesión como editor_lab desaparece la administración de usuarios y grupos. Géneros, personas y valoraciones se muestran en modo de consulta, mientras que películas permite agregar y modificar.",
            "06_editor_panel_restringido.jpg",
        ),
        (
            7,
            "Listado del editor sin eliminación",
            "La interfaz del editor permite trabajar con el catálogo dentro de su alcance. La ausencia de permiso de eliminación se valida adicionalmente mediante una prueba HTTP 403.",
            "07_editor_listado_sin_eliminar.jpg",
        ),
        (
            8,
            "Catálogo público",
            "La ruta principal presenta las diez películas cargadas, sus géneros, director, promedio y cantidad de valoraciones. Cada tarjeta permite consultar recomendaciones.",
            "08_catalogo_publico.jpg",
        ),
        (
            9,
            "Recomendaciones por género y puntuación",
            "La vista final excluye la película seleccionada, conserva las obras con géneros compartidos y muestra primero el mayor promedio disponible.",
            "09_recomendaciones_mismo_genero.jpg",
        ),
    ]
    for item in evidence_items:
        add_evidence_image(document, item[0], item[1], item[2], evidence_dir / item[3])

    document.add_page_break()
    document.add_paragraph("5. Reproducción y entrega", style="GLAB Heading 1")
    add_body(
        document,
        "El archivo README.md contiene la instalación completa, los comandos de migración, la preparación de usuarios, la ejecución en el puerto 8095, las pruebas y el procedimiento para configurar contraseñas sin almacenarlas en el repositorio.",
    )
    add_body(
        document,
        "El repositorio local queda preparado para asociarse al repositorio del equipo. La publicación remota requiere la URL exacta y una cuenta autorizada; por esa razón no se inventó un destino ni se expusieron credenciales.",
    )
    add_body(
        document,
        "La asistencia de inteligencia artificial utilizada durante la elaboración está declarada en PROMPTS.md. El proyecto conserva código legible, pruebas y datos reproducibles para que el estudiante pueda modificarlo y explicar cada decisión.",
    )
    add_status_box(
        document,
        "Entrega preparada",
        "código fuente, base local con datos, guía resuelta, capturas, documentación, pruebas y paquete ZIP revisable.",
    )

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("project_root", type=Path)
    args = parser.parse_args()
    build_document(args.source, args.destination, args.project_root)


if __name__ == "__main__":
    main()
